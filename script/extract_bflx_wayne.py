#!/usr/bin/env python3
"""Extract Wayne's BFLX cached bid, term, and IRR into tests/fixtures/bflx_wayne.json.

The workbook is not committed. On the RLMS box it lives at
crm/bids/bflx/01M5AYPAPH3Q40QKP0D9K8HQ58/. A copy uploaded with the build
spec is accepted as a fallback. Cached formula values are required
(data_only); this script does not reimplement the workbook.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PATHS = [
    ROOT / "crm/bids/bflx/01M5AYPAPH3Q40QKP0D9K8HQ58/BFLX Oct 2026 Portfolio Offering Pricing Model (BID).xlsx",
    Path("/home/ubuntu/.cursor/projects/workspace/uploads/BFLX_Oct_2026_Portfolio_Offering_Pricing_Model__BID__9551.xlsx"),
]
OUT = ROOT / "tests/fixtures/bflx_wayne.json"


def workbook_path() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    for candidate in DEFAULT_PATHS:
        if candidate.exists():
            return candidate
    raise SystemExit("BFLX workbook not found. Pass the xlsx path.")


def iso(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return None


def main() -> None:
    path = workbook_path()
    book = load_workbook(path, data_only=True)
    summary = book["Summary"]
    pricing = book["Pricing Summary"]
    by_deal: dict[str, dict[str, object]] = {}
    for row in range(2, 25):
        deal = summary.cell(row, 1).value
        if deal is None:
            continue
        by_deal[str(deal)] = {
            "lessee": summary.cell(row, 2).value,
            "company": summary.cell(row, 4).value,
            "rating": summary.cell(row, 5).value,
            "country": summary.cell(row, 6).value,
            "equipment": summary.cell(row, 8).value,
            "aar": summary.cell(row, 9).value,
            "commodity": summary.cell(row, 11).value,
        }
    deals = []
    for col in range(2, 23):
        deal = str(pricing.cell(7, col).value)
        meta = by_deal.get(deal, {})
        expiry = iso(pricing.cell(12, col).value)
        deals.append({
            "deal": deal,
            "cars": pricing.cell(9, col).value,
            "rent": pricing.cell(27, col).value,
            "rv": pricing.cell(31, col).value,
            "expiry": expiry,
            "buildYear": pricing.cell(15, col).value,
            "n": pricing.cell(89, col).value,
            "leaseEnd": iso(pricing.cell(88, col).value),
            "bidPerCar": pricing.cell(47, col).value,
            "acqPerCar": pricing.cell(50, col).value,
            "irr": pricing.cell(74, col).value,
            "balloon": pricing.cell(82, col).value,
            "leverage": pricing.cell(73, col).value,
            "debt": pricing.cell(41, col).value,
            "equity": pricing.cell(42, col).value,
            "company": meta.get("company"),
            "rating": meta.get("rating"),
            "aar": meta.get("aar"),
            "equipment": meta.get("equipment"),
        })
    cars = sum(row["cars"] or 0 for row in deals)
    bid = sum((row["bidPerCar"] or 0) * (row["cars"] or 0) for row in deals)
    payload = {
        "source": path.name,
        "settings": {
            "closeDate": "2026-07-31",
            "firstRentDate": "2026-08-01",
            "debtBase": 0.05,
            "spreadBps": 262,
            "equityReturn": 0.15,
            "debtClosePct": 0.005,
            "txnCostPct": 0.005,
            "leverageTarget": 0.8,
            "maintPcpm": 0,
            "usefulLifeYears": 50,
            "scrapPerCar": 7000,
            "discountRate": 0.09,
        },
        "deals": deals,
        "portfolio": {
            "cars": cars,
            "bid": bid,
            "note": "Portfolio IRR is recomputed from summed monthly cash flows in the golden test.",
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2) + "\n")
    print(f"wrote {OUT} ({len(deals)} deals, {cars} cars, bid {bid:.2f})")


if __name__ == "__main__":
    main()
